"""Regresión contra el solucionario del taller, no contra salidas de la IA."""
import json
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from decimal import Decimal
from PIL import Image
from django.test import TestCase, SimpleTestCase, override_settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import load_workbook
from pypdf import PdfReader
from .ciclo_contable import generar_asientos
from .reporte_utils import contexto_estados
from .models import AsientoContable, CuentaContable, Movimiento


def caso():
    return json.loads((Path(__file__).parent / 'testdata/caso_gaseosas.json').read_text())


def imagen():
    stream = BytesIO()
    Image.new('RGB', (500, 700), 'white').save(stream, format='PNG')
    return SimpleUploadedFile('caso.png', stream.getvalue(), content_type='image/png')


class ReglasCicloTests(SimpleTestCase):
    def test_cada_asiento_del_solucionario(self):
        asientos = generar_asientos(caso())
        esperado = [
            [('10','debe',10000),('50','haber',10000)],
            [('20','debe',4000),('10','haber',4000)],
            [('33','debe',3000),('10','haber',3000)],
            [('10','debe',3000),('12','debe',3000),('70','haber',6000)],
            [('10','debe',500),('20','haber',500)],
            [('10','debe',2700),('70','debe',300),('12','haber',3000)],
            [('20','debe',1000),('75','haber',1000)],
            [('20','debe',2000),('42','haber',2000)],
            [('10','debe',3000),('70','haber',3000)],
            [('66','debe',600),('20','haber',600)],
            [('62','debe',1200),('41','haber',1200)],
            [('63','debe',800),('63','debe',500),('10','haber',1300)],
            [('69','debe',3400),('20','haber',3400)],
            [('68','debe',135),('39','haber',135)],
        ]
        actual = [[(m['codigo'],m['tipo_movimiento'],m['monto']) for m in a['movimientos']] for a in asientos]
        self.assertEqual(actual, esperado)
        for a in asientos:
            self.assertEqual(sum(m['monto'] for m in a['movimientos'] if m['tipo_movimiento']=='debe'), sum(m['monto'] for m in a['movimientos'] if m['tipo_movimiento']=='haber'))
        self.assertEqual(sum(m['monto'] for a in asientos for m in a['movimientos'] if m['tipo_movimiento']=='debe'), Decimal('39135'))

    def test_datos_variados_no_son_un_solucionario_fijo(self):
        datos = caso()
        datos['operaciones'][0]['monto'] = '20000'
        datos['operaciones'][2]['monto'] = '6000'
        datos['operaciones'][-1]['monto'] = '2000'
        a = generar_asientos(datos)
        self.assertEqual(a[0]['movimientos'][0]['monto'], Decimal('20000'))
        self.assertEqual(a[-2]['movimientos'][0]['monto'], Decimal('3900'))
        self.assertEqual(a[-1]['movimientos'][0]['monto'], Decimal('270'))

    def test_cobro_fuera_de_plazo_sin_descuento(self):
        datos = caso();datos['operaciones'][5]['fecha']='2009-06-10'
        a = generar_asientos(datos)
        cobro = next(a for a in a if a['descripcion'].startswith('Por el cobro'))
        self.assertEqual(cobro['movimientos'][0]['monto'], Decimal('3000'))
        self.assertEqual(len(cobro['movimientos']),2)

    def test_datos_ambiguos_no_generan_asientos(self):
        cambios = [lambda d: d.update(observaciones=['Monto ilegible']),
                   lambda d: d['operaciones'].pop(),
                   lambda d: d['operaciones'][1].update(monto='NaN'),
                   lambda d: d['operaciones'][3].update(porcentaje_contado=110),
                   lambda d: d['operaciones'][5].update(fecha_factura='2009-05-16'),
                   lambda d: d['operaciones'][0].update(tipo='no_soportada'),
                   lambda d: d['operaciones'][2].update(residual_porcentaje=None)]
        for cambiar in cambios:
            datos=caso();cambiar(datos)
            with self.assertRaises(ValueError): generar_asientos(datos)

    @override_settings(GROQ_API_KEY='test-key')
    @patch('core.ai_diario.Groq')
    def test_vision_recibe_imagen_completa_y_dos_ampliaciones(self, groq):
        from .ai_diario import extraer_operaciones
        groq.return_value.chat.completions.create.return_value = SimpleNamespace(choices=[SimpleNamespace(finish_reason='stop',message=SimpleNamespace(content=json.dumps(caso())))])
        self.assertEqual(extraer_operaciones(imagen()), caso())
        params=groq.return_value.chat.completions.create.call_args.kwargs
        self.assertEqual(len(params['messages'][0]['content']),4)
        self.assertEqual(params['max_completion_tokens'],8192)
        groq.return_value.chat.completions.create.return_value.choices[0].finish_reason='length'
        with self.assertRaises(ValueError):extraer_operaciones(imagen())


class ImportacionReportesTests(TestCase):
    def importar(self):
        with patch('core.ai_diario.extraer_operaciones',return_value=caso()):
            r=self.client.post(reverse('cargar_imagen_diario'),{'imagen_caso':imagen(),'limpiar':'on'})
        self.assertRedirects(r,reverse('libro_diario'))

    def test_saldos_del_solucionario_y_clasificacion(self):
        self.importar();ctx=contexto_estados()
        esperado={'total_activo_corriente':13400,'total_activo_no_corriente':2865,'total_activos':16265,
                  'total_pasivos':3200,'total_patrimonio_con_resultados':13065,'resultados_acumulados':3065,
                  'total_ventas':8700,'total_costo_ventas':3400,'utilidad_bruta':5300,'total_gastos_operativos':2635,
                  'utilidad_operativa':2665,'total_otros_ingresos':1000,'total_otros_gastos':600,'gran_total_debe':39135}
        for clave,importe in esperado.items():self.assertEqual(ctx[clave],Decimal(importe),clave)
        self.assertTrue(ctx['esta_balanceado'])
        self.assertEqual(next(i['saldo'] for i in ctx['activo_no_corriente'] if i['cuenta'].codigo=='39'),Decimal('-135'))
        self.assertEqual(CuentaContable.objects.get(codigo='69').subcategoria,'costo_ventas')
        self.assertEqual(CuentaContable.objects.get(codigo='75').subcategoria,'otro_ingreso')
        self.assertEqual(CuentaContable.objects.get(codigo='66').subcategoria,'otro_gasto')
        for nombre in ('balance_general','estado_resultados'):
            r=self.client.get(reverse(nombre));self.assertEqual(r.status_code,200)
            self.assertContains(r,'S/ 3,065.00') if nombre=='estado_resultados' else self.assertContains(r,'S/ 16,265.00')

    def test_fallo_de_lectura_no_limpia_la_base(self):
        self.importar();ids=list(AsientoContable.objects.values_list('pk',flat=True))
        with patch('core.ai_diario.extraer_operaciones',side_effect=ValueError('Dato ilegible')):
            self.client.post(reverse('cargar_imagen_diario'),{'imagen_caso':imagen(),'limpiar':'on'})
        self.assertEqual(list(AsientoContable.objects.values_list('pk',flat=True)),ids)

    def test_fallo_al_guardar_revierte_limpieza_y_cuentas(self):
        self.importar()
        ids=list(AsientoContable.objects.values_list('pk',flat=True))
        conteo=Movimiento.objects.count()
        with patch('core.ai_diario.extraer_operaciones',return_value=caso()), patch('core.views.Movimiento.objects.bulk_create',side_effect=RuntimeError('Fallo de escritura')):
            r=self.client.post(reverse('cargar_imagen_diario'),{'imagen_caso':imagen(),'limpiar':'on'})
        self.assertRedirects(r,reverse('cargar_imagen_diario'))
        self.assertEqual(list(AsientoContable.objects.values_list('pk',flat=True)),ids)
        self.assertEqual(Movimiento.objects.count(),conteo)

    def test_cuenta_existente_corregida_al_reemplazar_caso(self):
        CuentaContable.objects.create(codigo='39',nombre='Depreciación acumulada',tipo='pasivo')
        self.importar();self.assertEqual(CuentaContable.objects.get(codigo='39').tipo,'activo')

    def test_reportes_con_saldos_correctos_y_autoria(self):
        self.importar()
        r=self.client.post(reverse('reporte_completo'),{'empresa':'ABC','formato':'pdf','accion':'descargar'})
        self.assertEqual(r.status_code,200)
        pdf=PdfReader(BytesIO(r.content));texto='\n'.join(p.extract_text() for p in pdf.pages)
        self.assertIn('BROADBROS',texto)
        self.assertNotIn('UNIVERSIDAD',texto)
        self.assertIn('Otros ingresos',texto);self.assertIn('Otros gastos',texto)
        self.assertIn('3065',texto.replace(',','').replace('.',''))
        r=self.client.post(reverse('reporte_completo'),{'empresa':'=1+1','formato':'excel','accion':'descargar'})
        wb=load_workbook(BytesIO(r.content),data_only=False)
        self.assertEqual(len(wb.sheetnames),5)
        self.assertEqual(wb.properties.creator,'BROADBROS')
        self.assertEqual(wb['Estado de Resultados']['B15'].value,3065)
        self.assertEqual(wb['Estado de Resultados']['A1'].data_type,'s')
        self.assertTrue(wb['Libro Diario'].freeze_panes)
        self.assertEqual(wb['Situación Financiera']['B'+str(wb['Situación Financiera'].max_row-1)].value,16265)
