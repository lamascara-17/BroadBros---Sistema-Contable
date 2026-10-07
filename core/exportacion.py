"""Libro de trabajo financiero con datos numéricos y formato de impresión."""
from io import BytesIO
from itertools import zip_longest
from math import ceil
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from .reporte_utils import contexto_estados

AZUL = '07393C'
LILA = 'E9F5F8'
NUMERO = '#,##0.00;[Red](#,##0.00);"—"'


def exportar_excel_contable(empresa='BroadBros'):
    ctx = contexto_estados()
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = 'BROADBROS'
    wb.properties.title = f'Informe contable · {empresa}'
    wb.properties.subject = 'Libros y estados financieros'

    def texto(cell, value):
        # Los textos ingresados por el usuario se guardan como texto, nunca fórmulas.
        cell.value = str(value)
        cell.data_type = 's'

    def hoja(titulo, columnas, anchos):
        ws = wb.create_sheet(titulo)
        ws.sheet_view.showGridLines = False
        for fila, valor in [(1, empresa), (2, titulo.upper()), (3, f"Al {ctx['fecha_cierre'].strftime('%d/%m/%Y')}" if ctx['fecha_cierre'] else 'Sin movimientos'), (4, 'Preparado por el sistema BROADBROS · Importes en soles')]:
            ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=len(columnas))
            texto(ws.cell(fila, 1), valor)
        ws.cell(1, 1).font = Font(name='Fira Mono', size=20, color=AZUL)
        ws.cell(2, 1).font = Font(name='Fira Mono', size=12, bold=True, color=AZUL)
        for fila in (3, 4):
            ws.cell(fila, 1).font = Font(name='Fira Mono', size=10, color='2C666E')
        ws.append([])
        for col, label in enumerate(columnas, 1):
            cell = ws.cell(6, col)
            texto(cell, label)
            cell.font = Font(name='Fira Mono', size=10, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor=AZUL)
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[6].height = 30
        ws.freeze_panes = 'A7'
        for i, ancho in enumerate(anchos, 1):
            ws.column_dimensions[get_column_letter(i)].width = ancho
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = '1:6'
        ws.oddFooter.left.text = 'BROADBROS'
        ws.oddFooter.right.text = 'Página &P de &N'
        ws.oddFooter.left.size = ws.oddFooter.right.size = 9
        return ws

    def fila(ws, valores, tipo='detalle'):
        index = ws.max_row + 1
        for col, valor in enumerate(valores, 1):
            cell = ws.cell(index, col)
            if valor is None:
                continue
            if isinstance(valor, str):
                texto(cell, valor)
            else:
                cell.value = valor
            cell.font = Font(name='Fira Mono', size=10, color='0A090C', bold=tipo != 'detalle')
            cell.alignment = Alignment(vertical='top', wrap_text=True, horizontal='right' if isinstance(valor, (int, float)) else 'left')
            if not isinstance(valor, str) and not hasattr(valor, 'year'):
                cell.number_format = NUMERO
                cell.alignment = Alignment(horizontal='right', vertical='top')
            if tipo == 'total':
                cell.fill = PatternFill('solid', fgColor=AZUL)
                cell.font = Font(name='Fira Mono', size=10, bold=True, color='FFFFFF')
            elif tipo in ('seccion', 'subtotal'):
                cell.fill = PatternFill('solid', fgColor=LILA)
                cell.border = Border(top=Side(style='thin', color='90DDF0'))
            elif index % 2:
                cell.fill = PatternFill('solid', fgColor='F0EDEE')
        ws.row_dimensions[index].height = 30
        return index

    def nota(ws, contenido):
        # Dividir las notas largas evita ocultar información en una fila fija.
        for inicio in range(0, len(contenido), 400):
            parte = contenido[inicio:inicio + 400]
            index = fila(ws, [parte])
            columnas = ws.max_column
            ws.merge_cells(start_row=index, start_column=1, end_row=index, end_column=columnas)
            ancho = sum(ws.column_dimensions[get_column_letter(c)].width for c in range(1, columnas + 1))
            ws.row_dimensions[index].height = max(30, 16 * (ceil(len(parte) / max(20, ancho * .8)) + 1))

    ws = hoja('Situación Financiera', ['Activo', 'Importe', 'Pasivo y patrimonio', 'Importe'], [55, 20, 55, 20])
    if ctx['nota_resultado']:
        nota(ws, ctx['nota_resultado'])
    for izquierda, derecha in ctx['filas_esf']:
        left = izquierda or ('', None, 'detalle')
        right = derecha or ('', None, 'detalle')
        index = fila(ws, [left[0], left[1], right[0], right[1]])
        for start, item in [(1, left), (3, right)]:
            if item[2] != 'detalle':
                for col in (start, start + 1):
                    cell = ws.cell(index, col)
                    cell.fill = PatternFill('solid', fgColor=AZUL if item[2] == 'total' else LILA)
                    cell.font = Font(name='Fira Mono', size=10, bold=True, color='FFFFFF' if item[2] == 'total' else AZUL)
    fila(ws, ['Activo = Pasivo + Patrimonio', ctx['total_activos'], 'Pasivo + Patrimonio', ctx['total_pasivo_patrimonio']], 'subtotal')
    fila(ws, ['Agrupación por código contable; revise vencimientos especiales.', None, None, None])

    ws = hoja('Estado de Resultados', ['Concepto', 'Importe (S/)'], [80, 24])
    if ctx['nota_resultado']:
        nota(ws, ctx['nota_resultado'])
    for concepto, saldo, tipo in ctx['filas_er']:
        fila(ws, [concepto, saldo], tipo)
    fila(ws, ['Solo se descuenta el impuesto a las ganancias registrado; no se presume una tasa tributaria.', None])

    ws = hoja('Libro Diario', ['Fecha', 'Asiento', 'Código', 'Cuenta', 'Glosa', 'Debe', 'Haber'], [14, 12, 10, 43, 66, 19, 19])
    for numero, asiento in enumerate(ctx['asientos'], 1):
        for mov in asiento.movimientos.all():
            index = fila(ws, [asiento.fecha, str(numero), mov.cuenta.codigo, mov.cuenta.nombre, asiento.descripcion,
                             mov.monto if mov.tipo == 'debe' else None, mov.monto if mov.tipo == 'haber' else None])
            ws.cell(index, 1).number_format = 'dd/mm/yyyy'
            ws.row_dimensions[index].height = 45
        fila(ws, [None, None, None, 'Total del asiento', None, asiento.total_debe, asiento.total_haber], 'subtotal')
    fila(ws, [None, None, None, 'TOTALES', None, ctx['gran_total_debe'], ctx['gran_total_haber']], 'total')

    ws = hoja('Libro Mayor', ['Código', 'Cuenta', 'Fecha', 'Referencia', 'Debe', 'Haber'], [10, 48, 14, 48, 19, 19])
    for dato in ctx['mayor_datos']:
        fila(ws, [dato['cuenta'].codigo, dato['cuenta'].nombre, None, None, None, None], 'seccion')
        for mov in dato['movimientos']:
            index = fila(ws, [None, None, mov.asiento.fecha, mov.asiento.descripcion,
                             mov.monto if mov.tipo == 'debe' else None, mov.monto if mov.tipo == 'haber' else None])
            ws.cell(index, 3).number_format = 'dd/mm/yyyy'
        fila(ws, [None, 'Sumas', None, None, dato['total_debe'], dato['total_haber']], 'subtotal')
        fila(ws, [None, 'Saldo final según naturaleza', None, None, dato['saldo_final'], None], 'subtotal')

    ws = hoja('Balance Comprobación', ['Código', 'Cuenta', 'Suma debe', 'Suma haber', 'Saldo deudor', 'Saldo acreedor'], [10, 55, 20, 20, 20, 20])
    for dato in ctx['bal_comp_datos']:
        fila(ws, [dato['cuenta'].codigo, dato['cuenta'].nombre, dato['total_debe'], dato['total_haber'], dato['saldo_deudor'], dato['saldo_acreedor']])
    fila(ws, [None, 'TOTALES', ctx['gran_total_debe'], ctx['gran_total_haber'], ctx['gran_saldo_deudor'], ctx['gran_saldo_acreedor']], 'total')
    for ws in wb:
        ws.print_area = f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
