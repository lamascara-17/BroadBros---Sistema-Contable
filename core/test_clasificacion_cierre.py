from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from django.test import TestCase, SimpleTestCase
from django.urls import reverse
from openpyxl import load_workbook
from pypdf import PdfReader
from .clasificacion import activo_no_corriente, fecha_cierre_del_caso
from .importacion_general import normalizar_borrador
from .models import AsientoContable
from .reporte_utils import contexto_estados
from .views import guardar_importacion


def pacifico():
    cuentas = {'101': ('Caja', 'activo'), '50': ('Capital', 'patrimonio'),
               '122': ('Equipo de Oficina', 'activo'), '465': ('Proveedor de equipos', 'pasivo'),
               '121': ('Clientes', 'activo'), '704': ('Servicios de diseño', 'ingreso'),
               '63': ('Alquiler', 'gasto')}
    operaciones = [(1, 'Aporte', [('101', 'debe', 20000), ('50', 'haber', 20000)]),
                   (3, 'Compra mixta', [('122', 'debe', 5000), ('101', 'haber', 2000), ('465', 'haber', 3000)]),
                   (8, 'Servicios', [('101', 'debe', 6000), ('121', 'debe', 3000), ('704', 'haber', 9000)]),
                   (15, 'Alquiler', [('63', 'debe', 1500), ('101', 'haber', 1500)]),
                   (20, 'Pago proveedor', [('465', 'debe', 1000), ('101', 'haber', 1000)]),
                   (28, 'Cobro cliente', [('101', 'debe', 2000), ('121', 'haber', 2000)])]
    return [{'fecha': f'2026-12-{d:02}', 'descripcion': glosa, 'movimientos': [
        {'codigo': c, 'nombre': cuentas[c][0], 'tipo_cuenta': cuentas[c][1], 'subcategoria': '',
         'tipo_movimiento': lado, 'monto': str(monto)} for c, lado, monto in lineas]} for d, glosa, lineas in operaciones]


class ClasificacionTests(SimpleTestCase):
    def test_naturaleza_y_eleccion_explicita(self):
        self.assertTrue(activo_no_corriente('122', 'Equipo de Oficina'))
        self.assertTrue(activo_no_corriente('2201', 'Mobiliario y Equipo de Oficina'))
        self.assertFalse(activo_no_corriente('122', 'Letras por cobrar'))
        self.assertFalse(activo_no_corriente('20', 'Equipos para la venta'))
        self.assertFalse(activo_no_corriente('122', 'Equipo de oficina', 'activo_corriente'))
        self.assertTrue(activo_no_corriente('123', 'Activo de largo plazo', 'activo_no_corriente'))

    def test_fecha_literal_prevalece_y_no_se_inventa(self):
        self.assertEqual(fecha_cierre_del_caso({'texto_leido': 'Fecha de cierre: 31/12/2026.', 'fecha_cierre': '2026-12-28'}), date(2026, 12, 31))
        self.assertIsNone(fecha_cierre_del_caso({'texto_leido': 'Sin cierre especificado'}))
        with self.assertRaises(ValueError): fecha_cierre_del_caso({'texto_leido': 'Fecha de cierre: 31/02/2026.'})


class PacificoReportesTests(TestCase):
    def test_importacion_y_los_tres_formatos_respetan_cierre_y_equipo(self):
        texto = '\n'.join(f"El {a['fecha'][8:]}/12/2026: {a['descripcion']}" for a in pacifico()) + '\nFecha de cierre: 31/12/2026.'
        borrador = normalizar_borrador({'asientos': pacifico(), 'pendientes': [], 'supuestos': []})
        with patch('core.importacion_general.proponer_caso_general', return_value=borrador):
            r = self.client.post(reverse('cargar_texto_diario'), {'texto_caso': texto})
        self.assertRedirects(r, reverse('libro_diario'))
        self.assertEqual(AsientoContable.objects.count(), 6)
        ctx = contexto_estados()
        self.assertEqual(ctx['fecha_cierre'], date(2026, 12, 31))
        self.assertEqual(ctx['asientos'][-1].fecha, date(2026, 12, 28))
        self.assertEqual([i['cuenta'].codigo for i in ctx['activo_no_corriente']], ['122'])
        self.assertEqual(ctx['total_activo_no_corriente'], Decimal('5000'))
        self.assertEqual(ctx['total_activo_corriente'], Decimal('24500'))
        self.assertEqual(ctx['total_activos'], Decimal('29500'))
        self.assertEqual(ctx['total_pasivos'], Decimal('2000'))
        self.assertEqual(ctx['total_patrimonio_con_resultados'], Decimal('27500'))
        self.assertEqual(ctx['resultado_neto'], Decimal('7500'))
        self.assertTrue(ctx['esta_balanceado'])
        self.assertContains(self.client.get(reverse('balance_general')), 'Al 31/12/2026')
        pdf = self.client.post(reverse('reporte_completo'), {'empresa': 'Servicios Pacífico S.A.C.', 'formato': 'pdf'})
        self.assertEqual(pdf.status_code, 200)
        texto_pdf = '\n'.join(p.extract_text() for p in PdfReader(BytesIO(pdf.content)).pages)
        self.assertIn('Al 31/12/2026', texto_pdf)
        self.assertIn('al 31/12/2026', texto_pdf)
        self.assertNotIn('Al 28/12/2026', texto_pdf)
        xlsx = self.client.post(reverse('reporte_completo'), {'formato': 'excel'})
        wb = load_workbook(BytesIO(xlsx.content), data_only=True)
        self.assertEqual(wb['Situación Financiera']['A3'].value, 'Al 31/12/2026')
        self.assertEqual(len(wb.sheetnames), 5)

    def test_cierre_anterior_a_operaciones_no_escribe_y_sin_cierre_usa_ultima_fecha(self):
        with self.assertRaises(ValueError): guardar_importacion(pacifico(), False, fecha_cierre=date(2026, 12, 1))
        self.assertEqual(AsientoContable.objects.count(), 0)
        guardar_importacion(pacifico(), False)
        self.assertEqual(contexto_estados()['fecha_cierre'], date(2026, 12, 28))
